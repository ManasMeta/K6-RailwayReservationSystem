import sys
from pathlib import Path

import openpyxl
import pytest

IMPORT_DIR = Path(__file__).resolve().parents[3] / "db" / "import"
sys.path.insert(0, str(IMPORT_DIR))

import load_train_details  # noqa: E402  (path is added just above)
import load_zones  # noqa: E402

SAMPLE = IMPORT_DIR.parent / "seeds" / "g2_train_details_sample.csv"


def test_read_rows_skips_header_and_keeps_raw_padding():
    rows = list(load_train_details.read_rows(SAMPLE))

    assert len(rows) == 1899
    assert rows[0][3] == "HWH "  # trimming happens in SQL, not Python


def test_read_rows_honours_limit():
    assert len(list(load_train_details.read_rows(SAMPLE, limit=5))) == 5


def test_read_rows_rejects_wrong_column_count(tmp_path):
    bad = tmp_path / "bad.csv"
    bad.write_text("a,b,c\n1,2,3\n", encoding="latin-1")

    with pytest.raises(ValueError, match="expected 12 columns"):
        list(load_train_details.read_rows(bad))


def test_zone_pairs_read_from_stations_sheet(tmp_path):
    workbook = openpyxl.Workbook()
    workbook.active.title = "Lookup_Tables"
    sheet = workbook.create_sheet("Stations")
    sheet.append(["Station Name", "Station Code", "Zone"])
    sheet.append(["New Delhi", "ndls", "nr"])
    sheet.append(["Howrah", "HWH", "ER"])
    sheet.append(["No Zone", "XXX", None])
    path = tmp_path / "book.xlsx"
    workbook.save(path)

    assert list(load_zones.read_pairs(path)) == [("NDLS", "NR"), ("HWH", "ER")]


def test_zone_sheet_without_zone_column_is_rejected(tmp_path):
    workbook = openpyxl.Workbook()
    workbook.active.title = "Stations"
    workbook.active.append(["Station Code", "Name"])
    path = tmp_path / "book.xlsx"
    workbook.save(path)

    with pytest.raises(ValueError, match="cannot find"):
        list(load_zones.read_pairs(path))
