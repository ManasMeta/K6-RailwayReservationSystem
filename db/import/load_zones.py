"""Load railway zones from the 'Stations' sheet of Indian_Railway_Reservation_System.xlsx.

Assumption (check against the real workbook): the sheet has a header row with
a column whose title contains "code" (station code) and one containing "zone".
Zone codes are stored upper-case; stations that are not in the database yet
are skipped (run load_train_details.py first). Reruns are no-ops.

Usage: python db/import/load_zones.py WORKBOOK.xlsx
"""
import os
import sys

import openpyxl
import psycopg


def find_columns(header):
    titles = [str(cell or "").strip().lower() for cell in header]
    code = next((i for i, t in enumerate(titles) if "code" in t and "zone" not in t), None)
    zone = next((i for i, t in enumerate(titles) if "zone" in t), None)
    if code is None or zone is None:
        raise ValueError(f"cannot find station-code and zone columns in {titles}")
    return code, zone


def read_pairs(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet_name = next(
        (name for name in workbook.sheetnames if "station" in name.lower()), None
    )
    if sheet_name is None:
        raise ValueError(f"no 'Stations' sheet in {workbook.sheetnames}")
    rows = workbook[sheet_name].iter_rows(values_only=True)
    code_col, zone_col = find_columns(next(rows))
    for row in rows:
        code, zone = row[code_col], row[zone_col]
        if code and zone:
            yield str(code).strip().upper(), str(zone).strip().upper()


def load(connection, pairs):
    pairs = list(pairs)
    with connection.cursor() as cursor:
        cursor.executemany(
            "INSERT INTO zones (zone_code, name) VALUES (%s, %s)"
            " ON CONFLICT (zone_code) DO NOTHING",
            [(zone, zone) for zone in sorted({z for _, z in pairs})],
        )
        cursor.executemany(
            "UPDATE stations SET zone_id = z.zone_id FROM zones z"
            " WHERE z.zone_code = %s AND stations.station_code = %s"
            " AND stations.zone_id IS NULL",
            [(zone, code) for code, zone in pairs],
        )
    connection.commit()
    return len(pairs)


def main(argv):
    if len(argv) != 1:
        sys.exit(__doc__)
    url = os.environ.get("DATABASE_URL")
    if not url:
        sys.exit("DATABASE_URL is not set")
    with psycopg.connect(url) as connection:
        count = load(connection, read_pairs(argv[0]))
    print(f"processed {count} station-zone rows")


if __name__ == "__main__":
    main(sys.argv[1:])
