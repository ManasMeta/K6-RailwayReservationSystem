"""Turn the 'PNR_Master' sheet of Indian_Railway_Reservation_System.xlsx into demo reservations (G4).

Every PNR is booked through the real engine: a reservations row, its passengers, then
CALL book_tickets(), so berths, RAC and WL come from G3's inventory exactly as for a live
booking. With --pay each booking is paid from the owner's wallet through pay_reservation()
(G5), and a PNR whose status says cancelled is then cancelled through cancel_passenger().
A PNR that is already in the database is skipped, so reruns are no-ops.

Assumptions (check against the real workbook): the sheet's first header row names its columns;
rows are one per passenger with the PNR repeated, or one per PNR with a passenger-count column.
Columns are found by header text, case-insensitively:

  required  "pnr", "train no", "journey date" (or "date"), "from", "to", "class"
  optional  "quota" (code or name, default GN), "passenger name", "age", "gender",
            "passengers" (count, when there are no name columns), "status", "email"

Workbook PNRs that are not 10 digits become "8" + their last 9 digits, so they stay stable.
Trains and stations missing from the database are reported and skipped. Booking-window rules
are switched off for this transaction only (demo dates may be in the past); --rebase-dates
instead moves the earliest journey to tomorrow and keeps the gaps between the rest.

Usage: python db/import/load_pnr_master.py WORKBOOK.xlsx [--user EMAIL] [--pay] [--rebase-dates]
--user owns rows without an "email" column value (for example alice@yatrio.test from
db/seeds/g1_demo_users.sql). Run load_train_details.py and generate_inventory.py first.
"""
import argparse
import os
import re
import sys
from collections import OrderedDict
from datetime import date, datetime, timedelta

import openpyxl
import psycopg

from load_workbook import _find, _int, _text, _title

COLUMNS = {
    "pnr": ((r"\bpnr\b",), (r"status",)),
    "train_no": ((r"train (no|number|num)", r"^train$"), (r"name",)),
    "run_date": ((r"journey date", r"travel date", r"date of journey", r"\bdoj\b", r"^date$"), ()),
    "from": ((r"^from", r"boarding", r"source"), ()),
    "to": ((r"^to\b", r"destination", r"alighting"), ()),
    "class_code": ((r"class",), ()),
    "quota": ((r"quota",), ()),
    "name": ((r"passenger name", r"^name$", r"^passenger$"), ()),
    "age": ((r"\bage\b",), ()),
    "gender": ((r"gender", r"\bsex\b"), ()),
    "count": ((r"no of passengers", r"passengers", r"\bpax\b"), (r"name",)),
    "status": ((r"status",), ()),
    "email": ((r"email",), ()),
}
REQUIRED = ("pnr", "train_no", "run_date", "from", "to", "class_code")
DATE_FORMATS = ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%b-%Y", "%d %b %Y")
CANCELLED = {"CAN", "CANCELLED", "CANCELED", "CNL"}


def normalise_pnr(value):
    digits = re.sub(r"\D", "", str(value or ""))
    if not digits:
        return None
    return digits if len(digits) == 10 else "8" + digits[-9:].zfill(9)


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    if text is None:
        return None
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"cannot read journey date {value!r}")


def parse_gender(value):
    text = (_text(value) or "").upper()
    return {"M": "Male", "F": "Female", "T": "Transgender"}.get(text[:1], "Male")


def find_columns(header):
    titles = [_title(c) for c in header]
    columns = {key: _find(titles, *patterns, exclude=exclude) for key, (patterns, exclude) in COLUMNS.items()}
    missing = [key for key in REQUIRED if columns[key] is None]
    if missing:
        raise ValueError(f"PNR_Master has no column for {', '.join(missing)} in {titles}")
    return columns


def read_bookings(path):
    """Group the PNR_Master rows into one booking per PNR, in sheet order."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = next((n for n in workbook.sheetnames if "pnr" in n.lower()), None)
    if sheet is None:
        raise ValueError(f"no 'PNR_Master' sheet in {workbook.sheetnames}")
    rows = workbook[sheet].iter_rows(values_only=True)
    columns = find_columns(next(rows))

    def cell(row, key):
        col = columns[key]
        return row[col] if col is not None and col < len(row) else None

    bookings = OrderedDict()
    for row in rows:
        pnr = normalise_pnr(cell(row, "pnr"))
        if pnr is None:
            continue
        booking = bookings.get(pnr)
        if booking is None:
            booking = bookings[pnr] = {
                "pnr": pnr,
                "train_no": re.sub(r"\D", "", str(cell(row, "train_no") or "")) or None,
                "run_date": parse_date(cell(row, "run_date")),
                "from": (_text(cell(row, "from")) or "").upper(),
                "to": (_text(cell(row, "to")) or "").upper(),
                "class_code": (_text(cell(row, "class_code")) or "").upper(),
                "quota": _text(cell(row, "quota")) or "GN",
                "status": (_text(cell(row, "status")) or "").upper(),
                "email": _text(cell(row, "email")),
                "passengers": [],
            }
        name = _text(cell(row, "name"))
        if name:
            booking["passengers"].append({
                "passenger_name": name[:100],
                "age": min(max(_int(cell(row, "age"), f"PNR {pnr} age") or 30, 1), 125),
                "gender": parse_gender(cell(row, "gender")),
            })
        elif not booking["passengers"]:
            count = min(max(_int(cell(row, "count"), f"PNR {pnr} passengers") or 1, 1), 6)
            booking["passengers"] = [
                {"passenger_name": f"Passenger {pnr[-4:]}-{i}", "age": 30, "gender": "Male"}
                for i in range(1, count + 1)
            ]
    return list(bookings.values())


def rebase_dates(bookings, today):
    """Move the earliest journey to tomorrow, keeping the gaps between the others."""
    dated = [b["run_date"] for b in bookings if b["run_date"]]
    if not dated:
        return bookings
    shift = (today + timedelta(days=1)) - min(dated)
    for booking in bookings:
        if booking["run_date"]:
            booking["run_date"] += shift
    return bookings


def _quota_code(cursor, quota):
    """A quota code, or a name such as "Tatkal" matched to the quota whose name starts with it."""
    row = cursor.execute(
        "SELECT quota_code FROM quotas WHERE quota_code = upper(%s) OR name ILIKE %s || '%%'"
        " ORDER BY quota_code = upper(%s) DESC, quota_code LIMIT 1",
        (quota, quota.split()[0], quota),
    ).fetchone()
    return row[0] if row else None


def _book(cursor, booking, user_id, pay):
    """Book one PNR; returns None on success or the reason it was skipped."""
    train = cursor.execute("SELECT train_id FROM trains WHERE train_no = %s", (booking["train_no"],)).fetchone()
    if train is None:
        return f"train {booking['train_no']} is not in the timetable"
    stops = {
        code: seq for code, seq in cursor.execute(
            "SELECT s.station_code, ts.seq FROM train_stops ts JOIN stations s USING (station_id)"
            " WHERE ts.train_id = %s AND (s.station_code = ANY(%s) OR upper(s.name) = ANY(%s))",
            (train[0], [booking["from"], booking["to"]], [booking["from"], booking["to"]]),
        ).fetchall()
    }
    names = dict(cursor.execute(
        "SELECT upper(s.name), s.station_code FROM stations s WHERE upper(s.name) = ANY(%s)",
        ([booking["from"], booking["to"]],),
    ).fetchall())
    from_code = booking["from"] if booking["from"] in stops else names.get(booking["from"])
    to_code = booking["to"] if booking["to"] in stops else names.get(booking["to"])
    if from_code not in stops or to_code not in stops or stops[from_code] >= stops[to_code]:
        return f"train {booking['train_no']} does not run {booking['from']} -> {booking['to']}"
    if cursor.execute("SELECT 1 FROM travel_classes WHERE class_code = %s", (booking["class_code"],)).fetchone() is None:
        return f"unknown class {booking['class_code']}"
    quota = _quota_code(cursor, booking["quota"])
    if quota is None:
        return f"unknown quota {booking['quota']}"
    if booking["run_date"] is None:
        return "no journey date"

    if cursor.execute("SELECT 1 FROM coaches WHERE train_id = %s LIMIT 1", (train[0],)).fetchone() is None:
        cursor.execute("SELECT generate_train_consist_coaches(%s)", (train[0],))
    reservation_id = cursor.execute(
        "INSERT INTO reservations (pnr, user_id, train_no, run_date, from_station_code, to_station_code,"
        " from_seq, to_seq, class_code, quota_code)"
        " VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s) RETURNING reservation_id",
        (booking["pnr"], user_id, booking["train_no"], booking["run_date"], from_code, to_code,
         stops[from_code], stops[to_code], booking["class_code"], quota),
    ).fetchone()[0]
    cursor.executemany(
        "INSERT INTO reservation_passengers (reservation_id, passenger_name, age, gender)"
        " VALUES (%s, %s, %s, %s)",
        [(reservation_id, p["passenger_name"], p["age"], p["gender"]) for p in booking["passengers"]],
    )
    cursor.execute("CALL book_tickets(%s)", (reservation_id,))
    if pay:
        cursor.execute("SELECT * FROM pay_reservation(%s, %s)", (user_id, booking["pnr"]))
        if booking["status"] in CANCELLED:
            cursor.execute(
                "SELECT cancel_passenger(passenger_id, 'Cancelled in PNR_Master')"
                " FROM reservation_passengers WHERE reservation_id = %s ORDER BY passenger_id",
                (reservation_id,),
            )
    return None


def load(connection, bookings, default_email=None, pay=False):
    """Book every PNR, each in its own savepoint; returns (booked, already there, skipped)."""
    booked, existing, skipped = [], [], []
    with connection.cursor() as cursor:
        cursor.execute("SELECT set_config('app.skip_booking_window_check', 'true', true)")
        cursor.execute("SELECT set_config('app.user_role', 'system', true)")
        for booking in bookings:
            if cursor.execute("SELECT 1 FROM reservations WHERE pnr = %s", (booking["pnr"],)).fetchone():
                existing.append(booking["pnr"])
                continue
            email = booking["email"] or default_email
            user = cursor.execute("SELECT user_id FROM users WHERE email = %s", (email,)).fetchone() if email else None
            if user is None:
                skipped.append((booking["pnr"], f"no user {email!r}; pass --user or add the account"))
                continue
            try:
                with connection.transaction():
                    reason = _book(cursor, booking, user[0], pay)
                    if reason:
                        raise LookupError(reason)
                booked.append(booking["pnr"])
            except (LookupError, psycopg.Error) as exc:
                skipped.append((booking["pnr"], str(exc).splitlines()[0]))
    return booked, existing, skipped


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("workbook")
    parser.add_argument("--user", help="email of the account that owns rows without an email")
    parser.add_argument("--pay", action="store_true", help="pay each booking from the owner's wallet")
    parser.add_argument("--rebase-dates", action="store_true", help="move the earliest journey to tomorrow")
    args = parser.parse_args(argv)

    url = os.environ.get("DATABASE_URL")
    if not url:
        sys.exit("DATABASE_URL is not set")
    bookings = read_bookings(args.workbook)
    if args.rebase_dates:
        rebase_dates(bookings, date.today())
    with psycopg.connect(url) as connection:
        booked, existing, skipped = load(connection, bookings, args.user, args.pay)
    print(f"{len(bookings)} PNRs read: {len(booked)} booked, {len(existing)} already loaded, {len(skipped)} skipped")
    for pnr, reason in skipped:
        print(f"  skipped {pnr}: {reason}")


if __name__ == "__main__":
    main()
